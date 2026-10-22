import re
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy.orm import scoped_session
from werkzeug.datastructures import FileStorage

from cndx_corp_back.common.decorators import component
from cndx_corp_back.common.log import CorpLogger
from cndx_corp_back.logic.building_vis import building_vis_constants as vis_constants
from cndx_corp_back.logic.building_vis.vis_file_check import VisFileCheck


@component
class VisEnergyDataCheck(VisFileCheck):
    def __init__(
        self,
        input_file: FileStorage,
        logger: CorpLogger,
        session: scoped_session,
        org_id: str,
        user_id: str,
        request_dict: dict[str, any],
    ) -> None:
        super().__init__(input_file, logger, session, org_id, user_id, request_dict)

    def execute(self):
        wb = openpyxl.load_workbook(self.input_file.stream, data_only=True)

        # 必須シートチェック
        self.check_required_sheet(wb, vis_constants.energy_required_sheet_names)
        if len(self.format_errors) > 1:
            # フォーマットチェックエラー
            self.logger.warning("フォーマットエラーのため、処理が中止されました。")
            self.logger.warning(f"{self.format_errors}")
            return self.format_errors

        ws_joho = wb["情報"]
        ws_bukken_joho = wb["物件情報"]
        ws_shubetsu_energy = wb["種別ごとエネルギー使用量"]
        ws_shubetsu_netsuryou = wb["種別ごと熱量換算値"]
        ws_shubetsu_co2_kiso = wb["種別ごとCO2基礎排出量"]
        ws_shubetsu_co2_choseigo = wb["種別ごとCO2調整後排出量"]
        ws_unit = wb["原単位分母"]
        ws_netsuryou_keisu = wb["熱量換算係数"]
        ws_co2_kiso_keisu = wb["CO2基礎排出係数"]
        ws_co2_choseigo_keisu = wb["CO2調整後排出係数"]

        setting_headers = [
            "企業名",
            "ダウンロード日時",
            "フォーマット種類",
            "フォーマットバージョン",
            "基準年月",
            "出力年数",
        ]
        self.check_required_headers(
            ws_joho, setting_headers, start_col="B", start_row=3, single_col_flg=True
        )

        building_headers = [
            "物件番号",
            "物件名",
            "ふりがな",
            "郵便番号",
            "国",
            "都道府県",
            "自治体コード",
            "市区町村番地",
            "建物名",
            "TEL",
            "備考",
            "延床面積",
        ]
        self.check_required_headers(
            ws_bukken_joho, building_headers, start_col="B", start_row=5
        )

        # 必須列チェック
        common_headers = [
            "物件番号",
            "物件名",
            "物件開始年月",
            "物件廃止年月",
            "エネルギー分類",
            "供給会社",
            "区分",
            "単位",
        ]
        self.check_required_headers(
            ws_shubetsu_energy, common_headers, start_col="B", start_row=5
        )
        self.check_required_headers(
            ws_shubetsu_netsuryou, common_headers, start_col="B", start_row=5
        )
        self.check_required_headers(
            ws_shubetsu_co2_kiso, common_headers, start_col="B", start_row=5
        )
        self.check_required_headers(
            ws_shubetsu_co2_choseigo, common_headers, start_col="B", start_row=5
        )
        self.check_required_headers(
            ws_netsuryou_keisu, common_headers, start_col="B", start_row=5
        )
        self.check_required_headers(
            ws_co2_kiso_keisu, common_headers, start_col="B", start_row=5
        )
        self.check_required_headers(
            ws_co2_choseigo_keisu, common_headers, start_col="B", start_row=5
        )

        unit_headers = [
            "物件番号",
            "物件名",
            "物件開始年月",
            "物件廃止年月",
            "原単位分母名",
            "単位",
        ]
        self.check_required_headers(ws_unit, unit_headers, start_col="B", start_row=5)

        if len(self.format_errors) > 1:
            # フォーマットチェックエラー
            self.logger.warning("フォーマットエラーのため、処理が中止されました。")
            self.logger.warning(f"{self.format_errors}")
            return self.format_errors

        self.check_column_length(
            ws_bukken_joho, col="B", length=100, start_row=6, not_fixed_check_flg=True
        )

        # 出力年数
        year_value = str(ws_joho["C8"].value)

        # 年数を抽出
        year_match = re.search(r"(\d+)年", year_value)
        years = int(year_match.group(1)) if year_match else None

        # 月数 = 年数 × 12
        months = years * 12 if years else None

        # 開始列のインデックス (N列 = 14)
        start_col_idx = column_index_from_string("N")

        # 最終列インデックス
        end_month_idx = start_col_idx + months - 1

        # 最終列名
        end_month_letter = get_column_letter(end_month_idx)

        # 種別ごとエネルギーチェック
        self.check_required_buildingNo(ws_unit, min_row=6)
        self.check_required_buildingNo(ws_shubetsu_energy, min_row=6)
        self.check_decimal_precision(
            ws_shubetsu_energy,
            start_col_letter="N",
            max_precision=25,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_shubetsu_energy,
            start_col_letter="N",
            header_row=5,
            max_col=end_month_idx,
        )

        self.check_header_month_format(
            ws_unit, start_col_letter="H", header_row=5, max_col=end_month_idx - 6
        )
        # 熱量換算値チェック
        self.check_required_buildingNo(ws_shubetsu_netsuryou, min_row=6)
        self.check_decimal_precision(
            ws_shubetsu_netsuryou,
            start_col_letter="N",
            max_precision=25,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_shubetsu_netsuryou,
            start_col_letter="N",
            header_row=5,
            max_col=end_month_idx,
        )
        # CO2基礎排出量チェック
        self.check_required_buildingNo(ws_shubetsu_co2_kiso, min_row=6)
        self.check_decimal_precision(
            ws_shubetsu_co2_kiso,
            start_col_letter="N",
            max_precision=18,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_shubetsu_co2_kiso,
            start_col_letter="N",
            header_row=5,
            max_col=end_month_idx,
        )
        # CO2調整後排出量チェック
        self.check_required_buildingNo(ws_shubetsu_co2_choseigo, min_row=6)
        self.check_decimal_precision(
            ws_shubetsu_co2_choseigo,
            start_col_letter="N",
            max_precision=18,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_shubetsu_co2_choseigo,
            start_col_letter="N",
            header_row=5,
            max_col=end_month_idx,
        )
        # 熱量換算係数チェック
        self.check_required_buildingNo(ws_netsuryou_keisu, min_row=6)
        self.check_decimal_precision(
            ws_netsuryou_keisu,
            start_col_letter="N",
            max_precision=14,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_netsuryou_keisu,
            start_col_letter="N",
            header_row=5,
            max_col=end_month_idx,
        )
        # CO2基礎排出係数チェック
        self.check_required_buildingNo(ws_co2_kiso_keisu, min_row=6)
        self.check_decimal_precision(
            ws_co2_kiso_keisu,
            start_col_letter="N",
            max_precision=14,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_co2_kiso_keisu, start_col_letter="N", header_row=5, max_col=end_month_idx
        )
        # CO2調整後
        # 係数チェック
        self.check_required_buildingNo(ws_co2_choseigo_keisu, min_row=6)
        self.check_decimal_precision(
            ws_co2_choseigo_keisu,
            start_col_letter="N",
            max_precision=14,
            max_scale=10,
            min_row=6,
            max_col_letter=end_month_letter,
        )
        self.check_header_month_format(
            ws_co2_choseigo_keisu,
            start_col_letter="N",
            header_row=5,
            max_col=end_month_idx,
        )

        if len(self.data_errors) > 1:
            # データチェックエラー
            self.logger.warning("データにエラーあるため、処理が中止されました。")
            self.logger.warning(f"{self.data_errors}")
            return self.data_errors

        # 拠点存在チェック
        location_cd_list = self.get_location_cd_list(
            ws_bukken_joho, start_row=6, start_col_letter="B"
        )
        self.check_building(ws_bukken_joho.title, location_cd_list)
        if len(self.building_error_list) > 1:
            # 拠点存在チェックエラー
            self.logger.warning("存在しない拠点があるため、処理が中止されました。")
            self.logger.warning(f"{self.building_error_list}")
            return self.building_error_list

        # 企業名チェック
        file_company_name = ws_joho["C3"].value
        self.check_company_name(file_company_name=file_company_name)

        if len(self.check_file_company_name) > 1:
            # フォーマットチェックエラー
            return self.check_file_company_name

        return file_company_name
